from io import BytesIO

from docx import Document
from pypdf import PdfReader


def extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    pages_text = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages_text.append(text)

    return "\n".join(pages_text).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    extracted_text = []

    # Normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            extracted_text.append(text)

    # Text inside tables
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()

                if text:
                    extracted_text.append(text)

    return "\n".join(extracted_text).strip()


def extract_txt_text(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8").strip()
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1").strip()


def extract_text(file_bytes: bytes, filename: str) -> str:
    filename = filename.lower()

    if filename.endswith(".pdf"):
        return extract_pdf_text(file_bytes)

    if filename.endswith(".docx"):
        return extract_docx_text(file_bytes)

    if filename.endswith(".txt"):
        return extract_txt_text(file_bytes)

    raise ValueError("Only PDF, DOCX, and TXT files are supported.")